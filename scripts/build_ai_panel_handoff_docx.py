#!/usr/bin/env python3
"""Build the English AI panel handoff guide as a readable Word document."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "AI_PANEL_HANDOFF_EN.md"
OUTPUT = ROOT / "docs" / "AI_PANEL_HANDOFF_EN.docx"
FONT_FAMILY = "Aptos"
MONO_FAMILY = "Aptos Mono"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_padding(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name: str = "Aptos", size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    if name in {"Aptos", "Aptos Display"}:
        name = FONT_FAMILY
    elif name == "Aptos Mono":
        name = MONO_FAMILY
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def remove_paragraph_borders(style_or_paragraph) -> None:
    """Remove inherited paragraph borders (notably Word's Title bottom rule)."""
    p_pr = style_or_paragraph._element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_inline(paragraph, text: str, size: float = 10.5) -> None:
    """Add a small Markdown subset while keeping code paths readable."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Aptos Mono", size=max(8.5, size - 0.5))
        else:
            label, _url = re.match(r"\[([^\]]+)\]\(([^\)]+)\)", token).groups()
            run = paragraph.add_run(label)
            set_run_font(run, size=size, color="1F4E79")
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Code Block"]
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.keep_together = True
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(run, name="Aptos Mono", size=8.2, color="2F2F2F")
        if index < len(lines) - 1:
            run.add_break()


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    # Make the header repeat when Word splits a long table across pages.
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    header_tr_pr.append(header_flag)
    for row_index, row in enumerate(rows):
        for col_index in range(column_count):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_padding(cell)
            text = row[col_index] if col_index < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, text, size=8.8 if row_index else 8.9)
            if row_index == 0:
                set_cell_shading(cell, "1F4E79")
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
            else:
                set_cell_shading(cell, "F4F7FA" if row_index % 2 == 0 else "FFFFFF")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(doc: Document, chinese: bool = False) -> None:
    global FONT_FAMILY, MONO_FAMILY
    FONT_FAMILY = "Hiragino Sans GB" if chinese else "Aptos"
    MONO_FAMILY = "Monaco" if chinese else "Aptos Mono"
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    remove_paragraph_borders(styles["Title"])
    normal = styles["Normal"]
    normal.font.name = FONT_FAMILY
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_FAMILY)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_FAMILY)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_FAMILY)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, space_before, space_after in (
        ("Heading 1", 16, 14, 7),
        ("Heading 2", 12.5, 10, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = styles[name]
        style.font.name = FONT_FAMILY
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_FAMILY)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_FAMILY)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_FAMILY)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.keep_with_next = True

    code = styles.add_style("Code Block", 1)
    code.font.name = MONO_FAMILY
    code._element.rPr.rFonts.set(qn("w:ascii"), MONO_FAMILY)
    code._element.rPr.rFonts.set(qn("w:hAnsi"), MONO_FAMILY)
    code._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FAMILY)
    code.font.size = Pt(8.2)
    code.paragraph_format.line_spacing = 1.0
    p_pr = code._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F4F7")
    p_pr.append(shd)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(3)
    footer_label = "AI 忠诚度面板交接  |  第 " if chinese else "AI Loyalty Panel Handoff  |  Page "
    run = footer.add_run(footer_label)
    set_run_font(run, size=8, color="6B7280")
    add_page_field(footer)
    if chinese:
        run = footer.add_run(" 页")
        set_run_font(run, size=8, color="6B7280")
    for run in footer.runs:
        set_run_font(run, size=8, color="6B7280")


def parse_markdown(doc: Document, markdown: str, chinese: bool = False) -> None:
    lines = markdown.splitlines()
    index = 0
    title_seen = False
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("```"):
            fence = line.strip()
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            add_code_block(doc, code_lines)
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            p = doc.add_paragraph(style="Title" if level == 1 and not title_seen else f"Heading {level}")
            if level == 1 and not title_seen:
                title_seen = True
                remove_paragraph_borders(p)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(text)
                set_run_font(run, name="Aptos Display", size=25, bold=True, color="000000")
                subtitle = doc.add_paragraph()
                subtitle.paragraph_format.space_after = Pt(14)
                subtitle_text = "中文业务流程、数据链路和部署交接说明" if chinese else "English operations, data-flow, and deployment handoff"
                r = subtitle.add_run(subtitle_text)
                set_run_font(r, size=11, color="4F6B82")
            else:
                add_inline(p, text, size=16 if level == 1 else 12.5 if level == 2 else 11)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[index + 1]):
            table_lines = [line]
            index += 1
            table_lines.append(lines[index])
            index += 1
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            rows = []
            for table_line in table_lines:
                cells = [cell.strip() for cell in table_line.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            continue
        list_match = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if list_match:
            style_name = "List Number" if list_match.group(2)[0].isdigit() else "List Bullet"
            p = doc.add_paragraph(style=style_name)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, list_match.group(3), size=10.2)
            index += 1
            continue
        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines) and lines[index].strip() and not re.match(r"^(#{1,3})\s+|^```|^\||^(\s*)([-*]|\d+\.)\s+", lines[index]):
            paragraph_lines.append(lines[index].strip())
            index += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines), size=10.5)


def main() -> None:
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else SOURCE
    output = Path(sys.argv[2]) if len(sys.argv) > 2 else OUTPUT
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    chinese = source.name.endswith("_ZH_REVIEW.md")
    configure_document(doc, chinese=chinese)
    parse_markdown(doc, source.read_text(encoding="utf-8"), chinese=chinese)
    properties = doc.core_properties
    properties.title = "AI Loyalty Panel Handoff Guide"
    properties.subject = "English operations and deployment handoff"
    properties.author = "AI Loyalty Engine"
    properties.comments = "Internal demo handoff document"
    doc.save(output)
    print(output)


if __name__ == "__main__":
    main()
