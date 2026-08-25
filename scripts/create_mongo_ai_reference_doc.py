from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/yangshikun/Documents/ChatGPT/Mogo Demo")
ARTIFACT_DIR = ROOT / "artifacts"
QA_DIR = ROOT / ".docx-qa" / "mongo-ai-reference"
OUTPUT = ARTIFACT_DIR / "MongoDB_AI参考系统_功能与架构分析.docx"
DIAGRAM = QA_DIR / "mongo_ai_reference_architecture.png"

SKILL_SCRIPTS = Path(
    "/Users/yangshikun/.codex/plugins/cache/openai-primary-runtime/documents/"
    "26.813.12317/skills/documents/scripts"
)
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry  # noqa: E402


# Resolved preset: standard_business_brief.
# Named visual override: "Mongo accent" (#00A35C / #E8F8F0), used only for
# the cover kicker, lead callouts, diagram, and selected status emphasis.
PAGE_WIDTH = Inches(8.5)
PAGE_HEIGHT = Inches(11)
MARGIN = Inches(1)
CONTENT_DXA = 9360
TABLE_INDENT = 120
CELL_MARGINS = {"top": 100, "bottom": 100, "start": 120, "end": 120}

FONT = "Arial Unicode MS"
LATIN_FONT = "Arial"
NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MONGO = "00A35C"
MONGO_DARK = "00684A"
MINT = "E8F8F0"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
GRAY = "667085"
DARK = "1F2937"
WHITE = "FFFFFF"
GOLD = "9A6700"
RED = "9B1C1C"
BORDER = "D7DEE7"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size: float | None = None, color: str = DARK,
                 bold: bool | None = None, italic: bool | None = None,
                 font_name: str = FONT):
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def style_paragraph(paragraph, *, before=0, after=6, line=1.10,
                    alignment=None, keep_with_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_with_next
    if alignment is not None:
        paragraph.alignment = alignment


def add_text(doc, text: str, *, size=11, color=DARK, bold=False,
             italic=False, before=0, after=6, line=1.10,
             alignment=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=False):
    p = doc.add_paragraph()
    style_paragraph(
        p,
        before=before,
        after=after,
        line=line,
        alignment=alignment,
        keep_with_next=keep_with_next,
    )
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, bold=True)
    return p


def add_bullet(doc, text: str, *, level=0, size=10.8, after=5, line=1.167):
    style_name = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def add_numbered(doc, text: str):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=10.8)
    return p


def add_callout(doc, title: str, body: str, *, fill=MINT, accent=MONGO_DARK):
    p = doc.add_paragraph()
    style_paragraph(p, before=2, after=10, line=1.12)
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8" if edge == "left" else "4")
        node.set(qn("w:space"), "5")
        node.set(qn("w:color"), accent if edge == "left" else BORDER)
        borders.append(node)
    p_pr.append(borders)
    r = p.add_run(title)
    set_run_font(r, size=10.8, color=accent, bold=True)
    r.add_break()
    r2 = p.add_run(body)
    set_run_font(r2, size=10.5, color=DARK)
    return p


def set_cell_text(cell, text: str, *, bold=False, color=DARK, size=9.7,
                  alignment=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    p.clear()
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold)


def add_matrix(doc, headers, rows, widths, *, header_fill=LIGHT_BLUE,
               font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, header_fill)
        set_cell_border(cell)
        set_cell_text(cell, header, bold=True, color=NAVY, size=9.3,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            set_cell_border(cell)
            if len(table.rows) % 2 == 1:
                shade_cell(cell, "FAFBFC")
            align = WD_ALIGN_PARAGRAPH.LEFT
            if idx == 0 and len(headers) > 2:
                align = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cell, str(value), size=font_size, alignment=align)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=CONTENT_DXA,
        indent_dxa=TABLE_INDENT,
        cell_margins_dxa=CELL_MARGINS,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_hyperlink(paragraph, text: str, url: str, color=BLUE):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), FONT)
    rpr.append(rfonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=8.5, color=GRAY)


def draw_wrapped_text(draw, box, text, font, fill, *, max_chars=13,
                      line_gap=8, align="center"):
    words = []
    for raw_line in text.split("\n"):
        if not raw_line:
            words.append("")
            continue
        current = ""
        for ch in raw_line:
            if len(current) >= max_chars:
                words.append(current)
                current = ch
            else:
                current += ch
        if current:
            words.append(current)
    bboxes = [draw.textbbox((0, 0), line, font=font) for line in words]
    heights = [b[3] - b[1] for b in bboxes]
    total_h = sum(heights) + max(0, len(words) - 1) * line_gap
    y = box[1] + (box[3] - box[1] - total_h) / 2
    for line, bb, h in zip(words, bboxes, heights):
        w = bb[2] - bb[0]
        if align == "center":
            x = box[0] + (box[2] - box[0] - w) / 2
        else:
            x = box[0] + 22
        draw.text((x, y), line, font=font, fill=fill)
        y += h + line_gap


def arrow(draw, x1, y1, x2, y2, color="#637083", width=8):
    draw.line((x1, y1, x2, y2), fill=color, width=width)
    size = 20
    draw.polygon(
        [(x2, y2), (x2 - size, y2 - size // 2), (x2 - size, y2 + size // 2)],
        fill=color,
    )


def create_architecture_diagram(path: Path):
    width, height = 2000, 1250
    img = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    font_path = "/System/Library/Fonts/Hiragino Sans GB.ttc"
    title_font = ImageFont.truetype(font_path, 54)
    subtitle_font = ImageFont.truetype(font_path, 27)
    col_font = ImageFont.truetype(font_path, 31)
    box_font = ImageFont.truetype(font_path, 27)
    note_font = ImageFont.truetype(font_path, 24)

    draw.text((80, 45), "MongoDB AI 参考系统架构", font=title_font, fill="#17324D")
    draw.text(
        (80, 112),
        "基于公开页面、接口响应与前端资源的结构推断",
        font=subtitle_font,
        fill="#667085",
    )

    columns = [
        (70, 185, 430, 930, "业务数据域", "#F5F7FA", "#667085"),
        (545, 185, 905, 930, "MongoDB 数据层", "#E8F8F0", "#00684A"),
        (1020, 185, 1380, 930, "API 与智能层", "#E8EEF5", "#1F4D78"),
        (1495, 185, 1855, 930, "运营体验层", "#F4F0FF", "#6941C6"),
    ]
    items = [
        ["会员 / CRM", "桌台与 Floor System", "POS / 酒店 / 活动", "治理规则与人工配置"],
        ["patron_profiles", "table_sessions / table_state", "offer_catalog / recommendations", "risk_cases / alerts / audit"],
        ["Python API（Uvicorn）", "确定性规则 / MQL", "LangGraph Agent 工作流", "Offer · Risk · Loss · Min-Bet"],
        ["Next.js 管理控制台", "Patron Eyes", "Offer Catalog / Insight", "Simulate / 人工审批"],
    ]

    for idx, (x1, y1, x2, y2, heading, fill, stroke) in enumerate(columns):
        draw.rounded_rectangle((x1, y1, x2, y2), radius=28, fill=fill, outline=stroke, width=4)
        draw.text((x1 + 28, y1 + 25), heading, font=col_font, fill=stroke)
        top = y1 + 105
        gap = 24
        bh = 125
        for item in items[idx]:
            box = (x1 + 25, top, x2 - 25, top + bh)
            draw.rounded_rectangle(box, radius=18, fill="#FFFFFF", outline=stroke, width=3)
            draw_wrapped_text(draw, box, item, box_font, stroke, max_chars=17)
            top += bh + gap

    arrow(draw, 438, 555, 532, 555)
    arrow(draw, 913, 555, 1007, 555)
    arrow(draw, 1388, 555, 1482, 555)

    loop_box = (280, 1010, 1720, 1135)
    draw.rounded_rectangle(loop_box, radius=24, fill="#FFF8E7", outline="#9A6700", width=4)
    draw_wrapped_text(
        draw,
        loop_box,
        "Human-in-the-loop：批准 / 拒绝 / 应用 / 发送 → 审计记录写回 MongoDB",
        col_font,
        "#7A5A00",
        max_chars=42,
    )
    draw.line((1675, 1008, 1675, 960, 725, 960, 725, 1008), fill="#9A6700", width=6)
    draw.polygon([(725, 1008), (712, 985), (738, 985)], fill="#9A6700")
    draw.text(
        (80, 1182),
        "说明：数据采集层的具体实现方式无法从公开资料确认。",
        font=note_font,
        fill="#667085",
    )

    img.save(path, quality=95, dpi=(180, 180))


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.right_margin = MARGIN
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(DARK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        1: (16, BLUE, 16, 8),
        2: (13, BLUE, 12, 6),
        3: (12, DARK_BLUE, 8, 4),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(10.8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.167

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("MONGODB AI REFERENCE SYSTEM REVIEW")
    set_run_font(hr, size=8.2, color=GRAY, bold=True, font_name=LATIN_FONT)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    add_page_number(fp)


def build_document():
    ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    QA_DIR.mkdir(parents=True, exist_ok=True)
    create_architecture_diagram(DIAGRAM)

    doc = Document()
    configure_document(doc)

    # Editorial cover pattern.
    add_text(
        doc,
        "REFERENCE SYSTEM REVIEW",
        size=10,
        color=MONGO_DARK,
        bold=True,
        before=72,
        after=18,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_text(
        doc,
        "MongoDB AI 参考系统",
        size=29,
        color=NAVY,
        bold=True,
        after=7,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_text(
        doc,
        "Casino Patron Marketing Platform\n功能与架构梳理",
        size=16,
        color=DARK_BLUE,
        bold=False,
        after=28,
        line=1.15,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_text(
        doc,
        "面向赌场营销、客户运营与风险治理的 AI Agent 应用参考",
        size=11.5,
        color=GRAY,
        italic=True,
        after=86,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_text(
        doc,
        "分析日期：2026 年 8 月 18 日",
        size=10.5,
        color=NAVY,
        bold=True,
        after=5,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    p_source = add_text(
        doc,
        "参考系统：",
        size=10,
        color=GRAY,
        after=24,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_hyperlink(p_source, "访问参考系统", "https://platform.ilovemongo.com/")
    add_callout(
        doc,
        "范围说明",
        "本报告介绍 MongoDB 同事提供的 AI 页面参考系统，重点梳理其业务功能、AI Agent 工作流、数据对象、治理机制与系统架构。",
        fill="F3FAF6",
        accent=MONGO_DARK,
    )
    doc.add_page_break()

    add_heading(doc, "执行摘要", 1)
    add_callout(
        doc,
        "一句话定位",
        "这是一个基于 MongoDB 运营数据构建的赌场智能营销与运营平台，通过多个 AI Agent 理解客户与桌台状态，生成优惠、风险和运营建议，并以人工审批与审计机制完成业务闭环。",
    )
    add_text(
        doc,
        "该系统的重点并非数据集成，而是展示数据进入 MongoDB 后，AI 应用能够创造的业务价值。平台把实时桌台运营、客户风险审核、优惠推荐、自然语言营销活动生成和最低投注额优化集中在同一管理控制台中。",
        after=8,
    )
    add_matrix(
        doc,
        ["观察快照", "当前公开页面数据", "业务含义"],
        [
            ["桌台", "30 张", "覆盖 A、B、C、VIP 四个区域"],
            ["在座客户", "293 人", "用于热力图、桌台钻取与风险分析"],
            ["优惠", "15 个；12 Active、3 Rejected", "体现优惠治理和生命周期"],
            ["客户推荐", "40 条", "包含相关性、置信度与下一步动作"],
        ],
        [1700, 2500, 5160],
    )

    add_heading(doc, "1. 系统定位与业务目标", 1)
    add_text(
        doc,
        "平台定位为赌场营销与现场运营的 AI 管理控制台。它把客户画像、桌台 Session、优惠目录、行为事件和风险记录组织为可被业务人员直接操作的工作流，而不是停留在只读报表或聊天机器人层面。",
    )
    add_bullet(doc, "实时感知：观察桌台占用、客户在座状态、投注与行为标签变化。")
    add_bullet(doc, "智能决策：为客户营销、风险审核和桌台运营生成建议。")
    add_bullet(doc, "治理执行：关键动作进入批准、拒绝、应用或发送流程。")
    add_bullet(doc, "审计闭环：保存建议、理由、状态、人工决策和执行结果。")

    add_heading(doc, "2. 功能模块总览", 1)
    add_matrix(
        doc,
        ["模块", "核心功能", "主要 AI / 决策能力"],
        [
            ["Patron Eyes", "桌台热力图、桌台钻取、在座客户、占用率、最低投注额", "Loss Potential、Risk Review、Min-Bet Optimization"],
            ["Offer Catalog", "优惠目录、状态管理、客户推荐、推荐分数与置信度", "Offer Agent、优惠生成、批准与拒绝"],
            ["Patron Insight", "输入客户 ID，生成单客户优惠匹配", "客户画像与优惠目录匹配"],
            ["Simulate", "模拟客户入座、投注、筹码、标签和 Active 状态变化", "用于触发热力图、规则与 AI 场景演示"],
        ],
        [1800, 3900, 3660],
        font_size=9.1,
    )

    add_heading(doc, "3. 系统架构", 1)
    add_text(
        doc,
        "下图基于公开页面、公开 API 响应、前端资源和响应头进行整理。Next.js、Uvicorn、同源 /api 路由及 LangGraph 标识属于可观察事实；具体数据采集方式与 Python Web 框架属于有限推断。",
        color=GRAY,
        size=10.2,
        after=10,
    )
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_after = Pt(5)
    run = p_img.add_run()
    picture = run.add_picture(str(DIAGRAM), width=Inches(6.45))
    doc_pr = picture._inline.docPr
    doc_pr.set("name", "MongoDB AI 参考系统架构图")
    doc_pr.set("descr", "从业务数据域到 MongoDB、Python API 与 LangGraph Agent，再到 Next.js 运营控制台和人工治理闭环的架构图")
    add_text(
        doc,
        "图 1｜MongoDB AI 参考系统架构（基于公开信息推断）",
        size=9.2,
        color=GRAY,
        italic=True,
        after=10,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_heading(doc, "3.1 前端与访问层", 2)
    add_bullet(doc, "前端使用 Next.js 和 React，静态资源由 Turbopack 构建。")
    add_bullet(doc, "管理控制台与后端 API 使用同一个域名，通过 /api 路径路由到动态服务。")
    add_bullet(doc, "页面使用定时轮询刷新热力图，不是以 WebSocket 为主的实时推送。")

    add_heading(doc, "3.2 后端与数据层", 2)
    add_bullet(doc, "API 响应由 Uvicorn 提供，说明后端为 Python ASGI 服务；FastAPI 或 Starlette 的可能性较高，但公开信息不足以完全确认。")
    add_bullet(doc, "MongoDB 承载客户、Session、桌台状态、优惠、推荐、风险、告警与审计数据。")
    add_bullet(doc, "领域接口围绕桌台、客户、优惠、风险和客户经理效率进行组织。")

    add_heading(doc, "3.3 AI 与治理层", 2)
    add_bullet(doc, "页面明确显示 LangGraph Online，说明多个 Agent 很可能由工作流编排。")
    add_bullet(doc, "确定性规则与 MQL 用于风险判断、Alert 分析和动作约束。")
    add_bullet(doc, "人工批准、拒绝、应用和发送形成 Human-in-the-loop 治理闭环。")

    add_heading(doc, "4. 主要功能详解", 1)
    add_heading(doc, "4.1 Patron Eyes：现场运营态势", 2)
    add_text(doc, "Patron Eyes 以桌台热力图展示桌台数量、在座客户、占用率、热门桌台、区域热度、游戏类型、平均投注和最低投注额。点击桌台后，可查看当前在座客户及其等级、地区、ADT、积分、Session 投注、筹码估算和行为标签。")
    add_bullet(doc, "启动单客户 Risk Review。")
    add_bullet(doc, "启动整桌 Loss Potential Agent。")
    add_bullet(doc, "运行 Min-Bet Optimizer，并批准、应用或拒绝建议。")
    add_bullet(doc, "模拟一轮投注，对 Active Alert 规则执行 MQL 分析。")

    add_heading(doc, "4.2 Offer Catalog：优惠与推荐治理", 2)
    add_text(doc, "Offer Catalog 管理酒店、餐饮、专车、娱乐票券、积分和现金回赠等权益，并展示成本、优先级、状态、AI 创建标识和拒绝原因。客户推荐同时保留相关性、置信度、推荐原因和 Next Best Action。")
    add_bullet(doc, "推荐生命周期：Proposed → Approved → Sent → Accepted / Rejected。")
    add_bullet(doc, "AI 创建的优惠可以被人工拒绝，并保留拒绝记录。")
    add_bullet(doc, "Next Best Action 包括立即发送、客户经理跟进和酒店套餐组合。")

    add_heading(doc, "4.3 Patron Insight：单客户匹配", 2)
    add_text(doc, "运营人员输入 Patron ID 后，系统生成该客户最匹配的优惠。该功能突出单客户快速决策，但当前导航名称为 Patron Insight，页面主标题仍显示 Offer Catalog，存在命名不一致。")

    add_heading(doc, "4.4 Simulate：场景模拟", 2)
    add_text(doc, "模拟模块允许选择桌台和客户，修改累计投注、筹码估算、行为标签及 Active/Inactive 状态。它适合展示数据变化如何驱动热力图、风险规则和 AI 建议。")

    add_heading(doc, "5. AI Agent 与决策工作流", 1)
    add_matrix(
        doc,
        ["Agent / 能力", "输入", "输出与人工动作"],
        [
            ["Loss Potential Agent", "桌台与客户 Session、投注分布、行为标签", "潜在损失或行为信号；进入人工复核"],
            ["Risk Review Agent", "客户等级、ADT、积分、风险标签、当前投注与行为", "风险案例；管理员批准、拒绝或请求更多信息"],
            ["Offer Agent", "自然语言营销目标、客群条件、优惠目录", "优惠草稿、匹配量与推荐；人工批准或拒绝"],
            ["Min-Bet Optimizer", "占用率、客户数量、投注分布、当前最低投注额", "最低投注建议；人工应用或拒绝"],
        ],
        [2050, 3550, 3760],
        font_size=9.0,
    )

    add_heading(doc, "5.1 典型业务闭环", 2)
    for step in [
        "系统读取客户、Session、桌台、优惠与风险上下文。",
        "确定性规则先识别资格、限制、风险和必须拦截的动作。",
        "AI Agent 生成解释、推荐、预测或运营建议。",
        "业务人员批准、拒绝、修改或请求更多信息。",
        "执行结果和人工决策写回 MongoDB，形成审计记录。",
    ]:
        add_numbered(doc, step)

    add_heading(doc, "6. 主要数据对象", 1)
    add_matrix(
        doc,
        ["数据域", "代表性集合", "作用"],
        [
            ["实时摄取", "patron_table_sessions；table_state_snapshots / history", "客户在座、投注、筹码、桌台人数、占用率与状态"],
            ["客户 360", "patron_profiles；patron_activity_events；patron_interaction_history", "等级、ADT、积分、偏好、行为和互动历史"],
            ["AI / NBA", "offer_catalog；offer_recommendations；patron_analysis_reports；chat_sessions / messages", "优惠、推荐、分析报告、Agent 对话与证据引用"],
            ["治理与风险", "patron_risk_cases；patron_alerts；alert_rules；offer_approval_audit；table_minbet_audit", "风险、告警、规则、审批与动作审计"],
        ],
        [1750, 3950, 3660],
        font_size=8.9,
    )

    add_heading(doc, "7. 值得借鉴的设计", 1)
    add_bullet(doc, "从运营数据直接进入可执行的 AI 工作流，而不是只展示分析结果。")
    add_bullet(doc, "多个专业 Agent 分工明确，分别服务营销、风险和桌台运营。")
    add_bullet(doc, "事实、分数、AI 建议和人工操作出现在同一业务上下文中。")
    add_bullet(doc, "推荐、风险和 Min-Bet 建议都具备状态生命周期。")
    add_bullet(doc, "关键动作必须人工批准，并保留原因、状态和审计记录。")
    add_bullet(doc, "模拟功能能够快速构造现场演示故事，降低对真实实时数据的依赖。")

    add_heading(doc, "8. 观察到的限制与注意事项", 1)
    add_matrix(
        doc,
        ["问题", "观察", "建议"],
        [
            ["公开访问", "未登录即可读取热力图和部分客户运营信息", "正式环境必须增加 OAuth/JWT、RBAC、脱敏、限流和审计"],
            ["数据新鲜度", "页面轮询，但部分 refreshedAt / lastActionAt 停留在 2026 年 5 月", "演示时说明为样例数据；生产环境展示同步时间与新鲜度"],
            ["状态矛盾", "存在 Closed 桌台仍有 Active 客户，以及单桌 133 人等异常", "增加唯一性、活动状态、容量和引用完整性校验"],
            ["语言一致性", "英文框架、繁体优惠和简体说明混用", "完整支持简体、繁体和英文国际化"],
            ["风险边界", "Session 数据足够做行为和营销风险，但不足以直接形成 AML 结论", "将事实、规则结果和 AI 推断分离；AML 需额外交易数据"],
        ],
        [1550, 3550, 4260],
        font_size=8.8,
    )

    add_heading(doc, "9. 系统业务价值", 1)
    add_text(doc, "该系统把桌台运营、客户画像、优惠营销、风险治理和人工审批集中到同一业务上下文中，使 AI 输出能够直接进入可执行、可追踪的运营流程。")
    add_callout(
        doc,
        "端到端业务闭环",
        "客户入座 → Session 行为变化 → 页面捕捉高价值或风险信号 → AI Agent 解释原因并生成优惠、风险或桌台运营建议 → 业务人员批准、拒绝或应用 → MongoDB 保存结果与审计记录。",
        fill="F6F9FC",
        accent=DARK_BLUE,
    )
    add_bullet(doc, "实时桌台与 Session 数据让运营人员能够快速定位值得关注的客户。")
    add_bullet(doc, "Offer、Risk 和 Min-Bet Agent 分工明确，分别服务营销、风险和桌台运营。")
    add_bullet(doc, "状态生命周期与人工审批机制使 AI 建议具备治理边界。")
    add_bullet(doc, "建议、原因、人工决策与执行结果写回 MongoDB，形成完整审计闭环。")

    add_heading(doc, "10. 总结", 1)
    add_text(
        doc,
        "Casino Patron Marketing Platform 展示了 MongoDB 数据在赌场营销和运营场景中的完整应用方式：从客户与桌台状态出发，通过多个 AI Agent 生成优惠、风险和运营建议，再由人工审批与审计机制完成闭环。它最大的参考价值，是把 AI 放进真实业务流程，而不是把 AI 作为独立聊天窗口。",
        after=8,
    )
    add_text(
        doc,
        "对外介绍时，可以将它准确描述为“MongoDB AI 应用参考系统”：它展示了 MongoDB 数据如何进入营销、风险和运营决策流程，并通过 AI Agent、人工审批与审计机制形成业务闭环。",
        bold=True,
        color=NAVY,
        after=12,
    )

    add_heading(doc, "附录：公开接口能力分组", 1)
    add_matrix(
        doc,
        ["领域", "代表性接口"],
        [
            ["桌台运营", "/api/tables/heatmap；/api/tables/{id}/patrons；/analyze；/optimize-minbet；/simulate-round"],
            ["优惠营销", "/api/offers/dashboard；/generate；/agent-chat；/{offerId}/approve；/{offerId}/reject"],
            ["风险治理", "/api/patrons/{id}/risk-case；/api/risk-cases/{id}/admin-decision；/api/alerts；/api/alert-rules"],
            ["客户经理", "/api/pr-efficiency/metrics；/kpi-search；/api/patrons/{id}/interactions"],
        ],
        [1800, 7560],
        font_size=9.0,
    )

    source_heading = add_heading(doc, "资料来源与方法", 2)
    source_heading.paragraph_format.space_before = Pt(7)
    source_heading.paragraph_format.space_after = Pt(3)
    p = add_text(doc, "公开页面与 API：", size=9.1, color=GRAY, after=1)
    add_hyperlink(p, "Casino Patron Marketing Platform", "https://platform.ilovemongo.com/")
    add_bullet(doc, "页面功能通过只读浏览确认；未触发 Generate、Agent、模拟、批准、拒绝或其他写入动作。", size=9.4, after=1, line=1.0)
    add_bullet(doc, "架构结论以公开响应头、前端资源引用、同源 API 和页面标识为依据；无法从公开信息确认的内容均按推断表述。", size=9.4, after=0, line=1.0)

    # Core properties and output.
    props = doc.core_properties
    props.title = "MongoDB AI 参考系统：功能与架构梳理"
    props.subject = "Casino Patron Marketing Platform 参考分析"
    props.author = "AI Decision Workspace Project"
    props.keywords = "MongoDB, AI Agent, Casino Marketing, Patron Risk, Architecture"
    props.comments = "基于公开页面与接口的只读分析"

    doc.save(OUTPUT)
    print(OUTPUT)


def build_compact_document():
    """Create the concise two-page handout requested for the demo discussion."""
    ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    QA_DIR.mkdir(parents=True, exist_ok=True)
    create_architecture_diagram(DIAGRAM)

    doc = Document()
    configure_document(doc)

    # Page 1 — what the reference system does.
    add_text(
        doc,
        "MONGODB AI REFERENCE SYSTEM",
        size=9.5,
        color=MONGO_DARK,
        bold=True,
        before=12,
        after=8,
    )
    add_text(
        doc,
        "MongoDB AI 参考系统",
        size=25,
        color=NAVY,
        bold=True,
        after=3,
    )
    add_text(
        doc,
        "Casino Patron Marketing Platform｜两页功能与架构摘要",
        size=12.5,
        color=DARK_BLUE,
        after=14,
    )
    add_callout(
        doc,
        "一句话结论",
        "这是一个面向赌场营销与现场运营的 AI 决策工作台：MongoDB 承载客户、桌台、Session、优惠与风险数据，多个 AI Agent 生成建议，再由业务人员审批、执行并保留审计记录。",
    )

    add_heading(doc, "它主要展示什么", 1)
    add_matrix(
        doc,
        ["模块", "现场看到的能力", "业务价值"],
        [
            ["Patron Eyes", "桌台热力图、在座客户、投注与行为标签", "实时发现高价值客户、风险和桌台机会"],
            ["Offer Catalog", "优惠目录、AI 推荐、置信度、批准与拒绝", "把推荐变成可治理、可执行的营销动作"],
            ["Patron Insight", "按 Patron ID 匹配最适合的优惠", "帮助客户经理快速做单客决策"],
            ["Simulate", "模拟入座、投注、筹码与状态变化", "快速触发 AI 场景，便于现场演示"],
        ],
        [1750, 3900, 3710],
        font_size=9.3,
    )

    add_heading(doc, "典型业务流程", 1)
    add_callout(
        doc,
        "从信号到行动",
        "客户入座 → Session 行为变化 → 页面捕捉高价值或风险信号 → AI Agent 解释原因并给出优惠、风险或桌台运营建议 → 业务人员批准、拒绝或应用 → MongoDB 保存结果与审计记录。",
        fill="F4F7FB",
        accent=DARK_BLUE,
    )
    add_bullet(doc, "系统不是独立聊天窗口，而是把 AI 放进桌台、客户、优惠和风险工作流。", size=10.2, after=2)
    add_bullet(doc, "推荐、风险和最低投注额建议都保留状态、理由与人工决策。", size=10.2, after=2)
    add_bullet(doc, "模拟功能可快速制造状态变化，用于观察规则、热力图和 Agent 输出。", size=10.2, after=2)

    add_text(
        doc,
        "说明：本文仅介绍该参考系统本身；内容来自公开页面、接口响应与前端资源的只读观察。",
        size=9.2,
        color=GRAY,
        italic=True,
        before=5,
        after=0,
    )

    doc.add_page_break()

    # Page 2 — architecture and key design characteristics.
    add_heading(doc, "系统架构（基于公开信息整理）", 1)
    add_text(
        doc,
        "可观察到 Next.js 前端、同源 /api、Python/Uvicorn、MongoDB 与 LangGraph Agent 标识；具体采集方式和部分后端实现属于合理推断。",
        size=9.8,
        color=GRAY,
        after=7,
    )
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_after = Pt(3)
    run = p_img.add_run()
    picture = run.add_picture(str(DIAGRAM), width=Inches(6.35))
    picture._inline.docPr.set("name", "MongoDB AI 参考系统架构图")
    picture._inline.docPr.set(
        "descr",
        "业务数据进入 MongoDB，经 Python API、规则与 LangGraph Agent 处理，在 Next.js 控制台中由人工审批、执行和审计。",
    )
    add_text(
        doc,
        "图 1｜业务数据 → MongoDB → API／规则／AI Agent → 运营控制台 → 人工治理与审计",
        size=8.8,
        color=GRAY,
        italic=True,
        after=8,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_heading(doc, "系统的 4 个设计特点", 1)
    add_bullet(doc, "以桌台或实时 Session 作为入口，点击后进入客户画像与 AI 分析。", size=10.0, after=2, line=1.05)
    add_bullet(doc, "把事实、规则结果、AI 推断和人工动作放在同一上下文中，避免黑盒推荐。", size=10.0, after=2, line=1.05)
    add_bullet(doc, "所有建议都具备状态生命周期：建议、批准、拒绝、发送、接受／未接受。", size=10.0, after=2, line=1.05)
    add_bullet(doc, "MongoDB 同时承载业务上下文、Agent 输出、人工决策和审计记录，形成闭环。", size=10.0, after=3, line=1.05)

    add_callout(
        doc,
        "系统定位",
        "这是一个以 MongoDB 为数据与审计底座、以规则和多个 AI Agent 为决策层、以运营控制台和人工审批为执行层的赌场智能营销与运营系统。",
        fill="F3FAF6",
        accent=MONGO_DARK,
    )

    p_source = add_text(doc, "参考：", size=8.8, color=GRAY, after=0)
    add_hyperlink(p_source, "Casino Patron Marketing Platform", "https://platform.ilovemongo.com/")

    props = doc.core_properties
    props.title = "MongoDB AI 参考系统：两页功能与架构摘要"
    props.subject = "Casino Patron Marketing Platform 简要分析"
    props.author = "AI Decision Workspace Project"
    props.keywords = "MongoDB, AI Agent, Casino Marketing, Architecture"
    props.comments = "两页精简版；基于公开页面与接口的只读分析"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
